"""浏览器版接口层（core/bridge.py）的测试。

这一层是给 GitHub Pages 上的在线版用的：Pyodide 里没有 urllib、没有多线程，
所以打分被拆成「拼 Prompt → JS 发请求 → 解析回复」三步。
这里把三步都测到，保证在线版和 exe 版走的是同一套逻辑。
"""

from __future__ import annotations

import base64
import csv
import io
import json
import unittest

from core.bridge import (
    bootstrap,
    export_available,
    export_b64,
    export_fallback_csv,
    export_name,
    fail_grade,
    finish_grade,
    mock_grade,
    parse_chat,
    prepare_grade,
)
from core.rules import DEFAULT_RUBRIC

CHAT = """张三
2026年09月26日 20:15
老师举了个打电话的例子讲三次握手，我一下就懂了。以前只知道背概念。

李四
2026年09月26日 20:16
老师举了个打电话的例子讲三次握手，我一下就懂了。以前只知道背概念。

王五
2026年09月26日 20:17
通过本次课程的学习，我深刻认识到团队合作具有重要意义。首先，我们要不断学习。
其次，我们要勇于创新。总而言之，这次课程让我受益匪浅，为我们今后的发展奠定基础。
"""

MODEL_REPLY = json.dumps({
    "criteria": [
        {"key": "understanding", "score": 4.5, "comment": "抓住了重点。"},
        {"key": "thinking", "score": 99, "comment": "超上限，应被夹住。"},
        {"key": "expression", "score": 1.6, "comment": "通顺。"},
    ],
    "overall_comment": "总评。",
    "risk_level": "none",
    "risk_reasons": [],
    "evidence": [],
}, ensure_ascii=False)


def jload(text: str):
    return json.loads(text)


class TestBootstrap(unittest.TestCase):
    def test_returns_default_rubric(self):
        data = jload(bootstrap())
        self.assertTrue(data["ok"])
        self.assertEqual(data["rubric"]["total"], 10)
        self.assertEqual(len(data["rubric"]["criteria"]), 3)
        self.assertIn("export_available", data)


class TestParseChat(unittest.TestCase):
    def test_parses_students(self):
        data = jload(parse_chat(json.dumps({"chat_text": CHAT}, ensure_ascii=False)))
        self.assertTrue(data["ok"])
        self.assertEqual([s["name"] for s in data["submissions"]], ["张三", "李四", "王五"])
        self.assertIn("stats", data)

    def test_empty_chat_is_rejected(self):
        data = jload(parse_chat(json.dumps({"chat_text": "   "})))
        self.assertFalse(data["ok"])
        self.assertIn("请先粘贴群聊记录", data["error"])

    def test_bad_json_is_rejected(self):
        data = jload(parse_chat("{不是 json"))
        self.assertFalse(data["ok"])

    def test_roster_mapping(self):
        payload = {"chat_text": CHAT, "roster_text": "姓名,学号,群昵称\n张三,2023123456,张三\n"}
        data = jload(parse_chat(json.dumps(payload, ensure_ascii=False)))
        self.assertEqual(data["submissions"][0]["student_id"], "2023123456")


class TestPrepareAndFinish(unittest.TestCase):
    def setUp(self):
        chat = jload(parse_chat(json.dumps({"chat_text": CHAT}, ensure_ascii=False)))
        self.submissions = chat["submissions"]
        self.rubric_json = json.dumps(DEFAULT_RUBRIC.to_dict(), ensure_ascii=False)

    def test_prepare_builds_prompts_and_similarity(self):
        data = jload(prepare_grade(json.dumps(self.submissions, ensure_ascii=False), self.rubric_json))
        self.assertTrue(data["ok"])
        self.assertEqual(data["total"], 3)
        first = data["items"][0]
        self.assertEqual(first["name"], "张三")
        self.assertEqual(first["max_total"], 10)
        self.assertEqual(len(first["messages"]), 2)
        self.assertEqual(first["messages"][0]["role"], "system")
        # 与李四雷同，必须被算出来并写进 Prompt
        self.assertTrue(first["similarity"])
        self.assertEqual(first["similarity"][0]["name"], "李四")
        self.assertIn("李四", first["messages"][1]["content"])
        self.assertIn("观点理解", first["messages"][1]["content"])

    def test_finish_parses_and_clamps(self):
        data = jload(prepare_grade(json.dumps(self.submissions, ensure_ascii=False), self.rubric_json))
        item = json.dumps(data["items"][0], ensure_ascii=False)
        out = jload(finish_grade(item, MODEL_REPLY, self.rubric_json))
        result = out["result"]
        self.assertEqual([c["score"] for c in result["criteria"]], [4.5, 3.0, 1.6])
        self.assertEqual(result["total"], 9.1)
        self.assertEqual(result["name"], "张三")
        self.assertEqual(result["source"], "llm")
        self.assertTrue(result["similarity"])

    def test_finish_handles_garbage_reply(self):
        data = jload(prepare_grade(json.dumps(self.submissions, ensure_ascii=False), self.rubric_json))
        item = json.dumps(data["items"][0], ensure_ascii=False)
        out = jload(finish_grade(item, "模型今天不想说话", self.rubric_json))
        self.assertTrue(out["result"]["error"])
        self.assertEqual(len(out["result"]["criteria"]), 3)

    def test_fail_grade_keeps_the_row(self):
        data = jload(prepare_grade(json.dumps(self.submissions, ensure_ascii=False), self.rubric_json))
        item = json.dumps(data["items"][1], ensure_ascii=False)
        out = jload(fail_grade(item, "API Key 被拒绝", self.rubric_json))
        result = out["result"]
        self.assertEqual(result["name"], "李四")
        self.assertIn("API Key", result["error"])
        self.assertEqual(result["total"], 0)


class TestMockGrade(unittest.TestCase):
    def test_mock_produces_results_for_everyone(self):
        chat = jload(parse_chat(json.dumps({"chat_text": CHAT}, ensure_ascii=False)))
        data = jload(mock_grade(json.dumps(chat["submissions"], ensure_ascii=False),
                                json.dumps(DEFAULT_RUBRIC.to_dict(), ensure_ascii=False)))
        results = data["results"]
        self.assertEqual(len(results), 3)
        self.assertTrue(all(r["source"] == "mock" for r in results))
        self.assertTrue(all(r["total"] > 0 for r in results))
        self.assertTrue(all("演示" in r["overall_comment"] for r in results))
        # 雷同依然要被标出来
        self.assertTrue(results[0]["similarity"])


class TestExport(unittest.TestCase):
    def setUp(self):
        chat = jload(parse_chat(json.dumps({"chat_text": CHAT}, ensure_ascii=False)))
        self.rubric_json = json.dumps(DEFAULT_RUBRIC.to_dict(), ensure_ascii=False)
        data = jload(mock_grade(json.dumps(chat["submissions"], ensure_ascii=False), self.rubric_json))
        self.results_json = json.dumps(data["results"], ensure_ascii=False)

    @unittest.skipUnless(export_available(), "需要 openpyxl")
    def test_xlsx_round_trip(self):
        from openpyxl import load_workbook

        b64 = export_b64(self.results_json, self.rubric_json, "计科2201", mock=True)
        blob = base64.b64decode(b64)
        self.assertEqual(blob[:2], b"PK")
        wb = load_workbook(io.BytesIO(blob))
        self.assertEqual(wb.sheetnames, ["成绩表", "评分规则", "雷同比对"])
        text = "\n".join(str(c.value) for row in wb["成绩表"].iter_rows() for c in row if c.value)
        self.assertIn("张三", text)
        self.assertIn("计科2201", text)

    def test_csv_fallback_contains_students(self):
        text = export_fallback_csv(self.results_json, self.rubric_json, mock=True)
        rows = list(csv.reader(io.StringIO(text)))
        self.assertIn("姓名", rows[0])
        flat = "\n".join(",".join(r) for r in rows)
        for name in ("张三", "李四", "王五"):
            self.assertIn(name, flat)
        self.assertIn("演示分数", flat)

    def test_export_name_is_safe(self):
        name = export_name(self.rubric_json, "计科/2201")
        self.assertTrue(name.endswith(".xlsx"))
        for bad in '<>:"/\\|?*':
            self.assertNotIn(bad, name)


class TestConsistencyWithServerPath(unittest.TestCase):
    """在线版和 exe 版必须给出同样的分数——这是不重写核心逻辑的意义所在。"""

    def test_mock_scores_match_grade_batch(self):
        from core.config import AppConfig
        from core.grader import grade_batch

        chat = jload(parse_chat(json.dumps({"chat_text": CHAT}, ensure_ascii=False)))
        subs = chat["submissions"]
        rubric_json = json.dumps(DEFAULT_RUBRIC.to_dict(), ensure_ascii=False)

        browser = jload(mock_grade(json.dumps(subs, ensure_ascii=False), rubric_json))["results"]
        server = [r.to_dict() for r in grade_batch(subs, DEFAULT_RUBRIC, AppConfig(api_key=""), mode="mock")]

        self.assertEqual([r["name"] for r in browser], [r["name"] for r in server])
        self.assertEqual([r["total"] for r in browser], [r["total"] for r in server])
        self.assertEqual([r["risk_level"] for r in browser], [r["risk_level"] for r in server])

    def test_prompt_matches_server_path(self):
        from core.grader import grade_one, chat_completion
        from core.config import AppConfig
        from unittest.mock import patch

        chat = jload(parse_chat(json.dumps({"chat_text": CHAT}, ensure_ascii=False)))
        subs = chat["submissions"][:1]
        rubric_json = json.dumps(DEFAULT_RUBRIC.to_dict(), ensure_ascii=False)
        browser_messages = jload(prepare_grade(json.dumps(subs, ensure_ascii=False),
                                               rubric_json))["items"][0]["messages"]

        captured = {}

        def fake(cfg, messages, **kwargs):
            captured["messages"] = messages
            return MODEL_REPLY

        with patch("core.grader.chat_completion", side_effect=fake):
            grade_one(rubric=DEFAULT_RUBRIC, cfg=AppConfig(api_key="sk-test1234567890"),
                      name=subs[0]["name"], content=subs[0]["content"],
                      student_id=subs[0].get("student_id"))

        self.assertEqual(
            [m["role"] for m in browser_messages],
            [m["role"] for m in captured["messages"]],
        )
        self.assertEqual(browser_messages[0]["content"], captured["messages"][0]["content"])
        self.assertEqual(browser_messages[1]["content"], captured["messages"][1]["content"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
