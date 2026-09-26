"""本地 HTTP 接口（core/webapp.py）的测试。

真起一个服务、真发请求——测的是老师实际会走的那条路径。
不需要 API Key：打分走演示模式。
"""

from __future__ import annotations

import io
import json
import random
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

from core.webapp import create_server, load_index_html, serve

CHAT = """张三
2026年09月26日 20:15
今天的课让我理解了三次握手的必要性。以前只知道背概念，现在明白了为什么是三次。

李四
2026年09月26日 20:20
我是李四，学号 2023123456。三权分立的核心在于制衡，这让我重新理解了制度设计。
"""

ROSTER = "姓名,学号,群昵称\n张三,2023123456,张三\n李四,2023123457,李四\n"


class ApiTestCase(unittest.TestCase):
    """起一个真实服务，供各用例发请求。"""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.config_path = Path(cls.tmp.name) / "config.ini"
        cls.server = create_server(random.randint(20000, 40000), config_path=cls.config_path)
        cls.base = f"http://127.0.0.1:{cls.server.state.port}"
        cls.thread = threading.Thread(target=serve, args=(cls.server,), daemon=True)
        cls.thread.start()
        time.sleep(0.25)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.tmp.cleanup()

    # ---- HTTP 小工具 ---------------------------------------------------- #

    def request(self, path: str, payload=None, method: str | None = None):
        url = self.base + path
        if payload is None and method != "POST":
            req = urllib.request.Request(url, method=method or "GET")
        else:
            req = urllib.request.Request(
                url,
                data=json.dumps(payload if payload is not None else {}).encode("utf-8"),
                headers={"Content-Type": "application/json"},
                method=method or "POST",
            )
        try:
            with urllib.request.urlopen(req, timeout=90) as resp:
                return resp.status, resp.read(), dict(resp.headers)
        except urllib.error.HTTPError as exc:
            return exc.code, exc.read(), dict(exc.headers)

    def get_json(self, path: str):
        status, body, _ = self.request(path)
        return status, json.loads(body)

    def post_json(self, path: str, payload=None):
        status, body, _ = self.request(path, payload if payload is not None else {})
        return status, json.loads(body)

    def parse_sample(self):
        _, data = self.post_json("/api/parse", {
            "chat_text": CHAT, "roster_text": ROSTER, "min_chars": 15,
        })
        return data["submissions"]

    def grade_and_wait(self, submissions, rubric, mode="mock", timeout=30.0):
        _, started = self.post_json("/api/grade", {
            "submissions": submissions, "rubric": rubric, "mode": mode,
        })
        job_id = started["job_id"]
        deadline = time.time() + timeout
        while time.time() < deadline:
            _, snap = self.get_json(f"/api/grade/status?job_id={job_id}")
            if snap["state"] in ("done", "error"):
                return snap
            time.sleep(0.15)
        self.fail("打分任务超时未完成")


class TestStaticAndHealth(ApiTestCase):
    def test_index_is_served(self):
        status, body, headers = self.request("/")
        self.assertEqual(status, 200)
        self.assertIn("text/html", headers["Content-Type"])
        self.assertIn("课堂心得评阅台", body.decode("utf-8"))

    def test_index_html_is_self_contained(self):
        html = load_index_html()
        self.assertNotIn("http://cdn", html)
        self.assertNotIn("https://cdn", html)
        self.assertNotIn("<script src=", html)     # 不依赖外部脚本
        self.assertIn("</html>", html)

    def test_health(self):
        status, data = self.get_json("/api/health")
        self.assertEqual(status, 200)
        self.assertTrue(data["ok"])

    def test_unknown_path_is_404_with_chinese_message(self):
        status, data = self.get_json("/api/nope")
        self.assertEqual(status, 404)
        self.assertIn("没有这个地址", data["error"])


class TestConfigApi(ApiTestCase):
    def test_default_config_is_mock_mode(self):
        _, data = self.get_json("/api/config")
        cfg = data["config"]
        self.assertFalse(cfg["has_api_key"])
        self.assertEqual(cfg["mode"], "mock")
        self.assertIn("config_path", cfg)

    def test_saving_api_key_switches_to_live(self):
        status, data = self.post_json("/api/config", {
            "api_key": "sk-abcdef1234567890", "model": "test-model", "class_name": "计科2201",
        })
        self.assertEqual(status, 200)
        self.assertTrue(data["config"]["has_api_key"])
        self.assertEqual(data["config"]["mode"], "live")
        self.assertEqual(data["config"]["model"], "test-model")
        self.assertTrue(self.config_path.exists())
        content = self.config_path.read_text(encoding="utf-8")
        self.assertIn("sk-abcdef1234567890", content)
        self.assertIn("计科2201", content)

    def test_masked_key_is_not_written_back(self):
        self.post_json("/api/config", {"api_key": "sk-realkey1234567890"})
        _, data = self.get_json("/api/config")
        masked = data["config"]["api_key_masked"]
        self.assertIn("*", masked)
        # 前端把脱敏 Key 回传时不能覆盖真实 Key
        self.post_json("/api/config", {"api_key": masked, "model": "m2"})
        raw = self.config_path.read_text(encoding="utf-8")
        self.assertIn("sk-realkey1234567890", raw)
        self.assertNotIn("*", raw)

    def test_concurrency_is_clamped(self):
        status, data = self.post_json("/api/config", {"concurrency": 999})
        self.assertEqual(status, 200)
        self.assertLessEqual(data["config"]["concurrency"], 16)


class TestRubricApi(ApiTestCase):
    def test_default_rubric(self):
        _, data = self.get_json("/api/rubric")
        rubric = data["rubric"]
        self.assertEqual(rubric["total"], 10)
        self.assertEqual([c["name"] for c in rubric["criteria"]],
                         ["观点理解", "个人思考", "文字表达"])

    def test_validate_reports_chinese_errors(self):
        _, data = self.post_json("/api/rubric/validate", {
            "rubric": {"name": "坏规则", "criteria": [{"key": "a", "name": "维度", "max_score": 0}]},
        })
        self.assertTrue(any("满分必须大于 0" in e for e in data["errors"]))

    def test_validate_accepts_default(self):
        _, rubric_data = self.get_json("/api/rubric")
        _, data = self.post_json("/api/rubric/validate", {"rubric": rubric_data["rubric"]})
        self.assertEqual(data["errors"], [])


class TestParseApi(ApiTestCase):
    def test_parse_extracts_students(self):
        subs = self.parse_sample()
        self.assertEqual([s["name"] for s in subs], ["张三", "李四"])
        self.assertEqual(subs[1]["student_id"], "2023123457")   # 以名单为准

    def test_parse_returns_stats_and_warnings(self):
        _, data = self.post_json("/api/parse", {"chat_text": CHAT})
        self.assertIn("messages", data["stats"])
        self.assertIsInstance(data["warnings"], list)

    def test_empty_chat_is_rejected(self):
        status, data = self.post_json("/api/parse", {"chat_text": "   "})
        self.assertEqual(status, 400)
        self.assertIn("请先粘贴群聊记录", data["error"])

    def test_invalid_json_is_rejected(self):
        req = urllib.request.Request(
            self.base + "/api/parse", data=b"not json",
            headers={"Content-Type": "application/json"}, method="POST",
        )
        try:
            urllib.request.urlopen(req, timeout=10)
            self.fail("应当返回 400")
        except urllib.error.HTTPError as exc:
            self.assertEqual(exc.code, 400)

    def test_exclude_names(self):
        _, rubric_data = self.get_json("/api/rubric")
        _, data = self.post_json("/api/parse", {
            "chat_text": "老师：本次心得不少于 200 字。\n张三：今天学到很多，收获很大。",
            "exclude_names": "老师",
        })
        self.assertEqual([s["name"] for s in data["submissions"]], ["张三"])


class TestGradeApi(ApiTestCase):
    def test_mock_grading_end_to_end(self):
        subs = self.parse_sample()
        _, rubric_data = self.get_json("/api/rubric")
        snap = self.grade_and_wait(subs, rubric_data["rubric"], mode="mock")
        self.assertEqual(snap["state"], "done")
        self.assertEqual(len(snap["results"]), 2)
        self.assertTrue(all(r["source"] == "mock" for r in snap["results"]))
        self.assertTrue(all(r["total"] > 0 for r in snap["results"]))
        self.assertTrue(all(len(r["criteria"]) == 3 for r in snap["results"]))

    def test_progress_is_reported(self):
        subs = self.parse_sample()
        _, rubric_data = self.get_json("/api/rubric")
        _, started = self.post_json("/api/grade", {
            "submissions": subs, "rubric": rubric_data["rubric"], "mode": "mock",
        })
        self.assertEqual(started["total"], 2)
        self.assertEqual(started["mode"], "mock")

    def test_live_mode_without_key_is_rejected(self):
        self.post_json("/api/config", {"api_key": ""})
        subs = self.parse_sample()
        _, rubric_data = self.get_json("/api/rubric")
        status, data = self.post_json("/api/grade", {
            "submissions": subs, "rubric": rubric_data["rubric"], "mode": "live",
        })
        self.assertEqual(status, 400)
        self.assertIn("API Key", data["error"])

    def test_empty_submissions_is_rejected(self):
        _, rubric_data = self.get_json("/api/rubric")
        status, data = self.post_json("/api/grade", {"submissions": [], "rubric": rubric_data["rubric"]})
        self.assertEqual(status, 400)
        self.assertIn("解析预览", data["error"])

    def test_bad_rubric_is_rejected(self):
        subs = self.parse_sample()
        status, data = self.post_json("/api/grade", {
            "submissions": subs,
            "rubric": {"name": "坏", "criteria": [{"key": "a", "name": "维度", "max_score": 0}]},
        })
        self.assertEqual(status, 400)
        self.assertIn("评分规则有问题", data["error"])

    def test_unknown_job_id(self):
        status, data = self.get_json("/api/grade/status?job_id=nope")
        self.assertEqual(status, 404)
        self.assertIn("找不到这个打分任务", data["error"])


class TestExportApi(ApiTestCase):
    def graded(self):
        subs = self.parse_sample()
        _, rubric_data = self.get_json("/api/rubric")
        snap = self.grade_and_wait(subs, rubric_data["rubric"], mode="mock")
        return snap["results"], rubric_data["rubric"]

    def test_export_returns_xlsx_attachment(self):
        results, rubric = self.graded()
        status, body, headers = self.request("/api/export", {
            "results": results, "rubric": rubric, "class_name": "计科2201", "mock": True,
        })
        self.assertEqual(status, 200)
        self.assertEqual(body[:2], b"PK")                       # zip/xlsx 魔数
        self.assertIn("attachment", headers["Content-Disposition"])
        wb = load_workbook(io.BytesIO(body))
        self.assertEqual(wb.sheetnames, ["成绩表", "评分规则", "雷同比对"])

    def test_download_filename_survives_http_headers(self):
        """中文文件名必须百分号编码后再放进 HTTP 头，否则浏览器拿到的是乱码。"""
        results, rubric = self.graded()
        _, _, headers = self.request("/api/export", {
            "results": results, "rubric": rubric, "class_name": "计科2201",
        })
        encoded = headers.get("X-Filename-Encoded")
        self.assertTrue(encoded, "应当带 X-Filename-Encoded")
        self.assertTrue(encoded.isascii(), "HTTP 头里必须是纯 ASCII")
        name = urllib.parse.unquote(encoded)
        self.assertIn("计科2201", name)
        self.assertTrue(name.endswith(".xlsx"))
        self.assertNotIn("\ufffd", name)
        # Content-Disposition 里也要有 RFC 5987 的中文名
        self.assertIn("filename*=UTF-8''", headers["Content-Disposition"])

    def test_edited_scores_are_exported(self):
        results, rubric = self.graded()
        results[0]["criteria"][0]["score"] = 1.0               # 老师手动改分
        status, body, _ = self.request("/api/export", {
            "results": results, "rubric": rubric,
        })
        self.assertEqual(status, 200)
        wb = load_workbook(io.BytesIO(body))
        ws = wb["成绩表"]
        values = [c.value for row in ws.iter_rows() for c in row]
        # 总分必须按改后的维度分重算（1.0 + 其它维度）
        expected = round(sum(c["score"] for c in results[0]["criteria"]), 2)
        self.assertIn(expected, values)

    def test_export_without_results_is_rejected(self):
        status, data = self.post_json("/api/export", {"results": []})
        self.assertEqual(status, 400)
        self.assertIn("没有可导出的成绩", data["error"])

    def test_export_with_name_only_result_works(self):
        """结果缺字段也不能崩——导出要尽量兜住。"""
        status, body, _ = self.request("/api/export", {"results": [{"name": "张三"}]})
        self.assertEqual(status, 200)
        self.assertEqual(body[:2], b"PK")


class TestJobSurvivesBrokenStdout(ApiTestCase):
    """回归测试：控制台编码不了中文时，打分任务也绝不能卡死。

    背景：任务线程里原本有一句中文 print 写在 try 之外。在 stdout 编码不了中文的
    环境（Windows 默认 ANSI 代码页、或输出被重定向到管道/文件）里，这句 print 会抛
    UnicodeEncodeError，线程当场死掉，任务永远停在 running——界面上就一直转圈，
    老师永远等不到结果，也看不到任何错误。

    这里把 sys.stdout 换成一个写不了中文的流来复现当时的场景。
    """

    def test_job_completes_with_unencodable_stdout(self):
        broken = io.TextIOWrapper(io.BytesIO(), encoding="cp1252", errors="strict")
        original = sys.stdout
        sys.stdout = broken
        try:
            # 先确认这个流真的写不了中文，否则这个用例就失去意义了
            with self.assertRaises(UnicodeEncodeError):
                broken.write("中文")
            subs = self.parse_sample()
            _, rubric_data = self.get_json("/api/rubric")
            snap = self.grade_and_wait(subs, rubric_data["rubric"], mode="mock", timeout=25)
        finally:
            sys.stdout = original

        self.assertEqual(snap["state"], "done", f"任务没有正常结束：{snap}")
        self.assertEqual(len(snap["results"]), 2)
        self.assertTrue(all(r["total"] > 0 for r in snap["results"]))


class TestShutdown(ApiTestCase):
    def test_shutdown_stops_the_server(self):
        # 单独起一个服务，避免影响其它用例
        server = create_server(random.randint(40000, 50000), config_path=self.config_path)
        base = f"http://127.0.0.1:{server.state.port}"
        thread = threading.Thread(target=serve, args=(server,), daemon=True)
        thread.start()
        time.sleep(0.2)

        req = urllib.request.Request(
            base + "/api/shutdown", data=b"{}",
            headers={"Content-Type": "application/json"}, method="POST",
        )
        with urllib.request.urlopen(req, timeout=10) as resp:
            self.assertEqual(resp.status, 200)

        thread.join(timeout=8)
        self.assertFalse(thread.is_alive(), "退出接口应当让服务停下来")
        self.assertTrue(server.state.should_exit)


if __name__ == "__main__":
    unittest.main(verbosity=2)
