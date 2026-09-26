"""Excel 导出（core/exporter.py）的测试。"""

from __future__ import annotations

import io
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from core.config import AppConfig
from core.exporter import save, suggested_filename, to_bytes
from core.grader import GradeResult, grade_batch
from core.rules import DEFAULT_RUBRIC

CONCRETE = (
    "老师提到 TCP 三次握手的时候举了个例子，说这就像打电话先确认对方在不在。"
    "我上次做实验抓包看到 SYN、ACK 那几个包，当时完全没看懂，这节课突然就明白了。"
)
COPIED = CONCRETE
UNRELATED = "今天食堂的饭特别好吃，我和室友打了一下午球，挺开心的。"


def mock_cfg() -> AppConfig:
    return AppConfig(api_key="")


def sample_results() -> list[GradeResult]:
    subs = [
        {"name": "张三", "content": CONCRETE, "student_id": "2023001", "raw_name": "赢赢"},
        {"name": "李四", "content": COPIED, "student_id": "2023002", "raw_name": "logic"},
        {"name": "王五", "content": UNRELATED, "student_id": "2023003", "raw_name": "王五"},
    ]
    results = grade_batch(subs, DEFAULT_RUBRIC, mock_cfg())
    results[0].manual_comment = "已与学生确认过。"
    results[0].reviewed = True
    return results


def load(data: bytes):
    return load_workbook(io.BytesIO(data))


class TestWorkbookStructure(unittest.TestCase):
    def setUp(self):
        self.wb = load(to_bytes(sample_results(), DEFAULT_RUBRIC, class_name="计科2201"))

    def test_three_sheets(self):
        self.assertEqual(self.wb.sheetnames, ["成绩表", "评分规则", "雷同比对"])

    def test_first_sheet_is_active(self):
        self.assertEqual(self.wb.active.title, "成绩表")

    def test_grade_sheet_has_all_students(self):
        ws = self.wb["成绩表"]
        text = "\n".join(
            str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
        )
        for name in ("张三", "李四", "王五"):
            self.assertIn(name, text)
        self.assertIn("2023001", text)
        self.assertIn("赢赢", text)          # 群昵称要保留，方便老师核对
        self.assertIn("计科2201", text)
        self.assertIn("总分", text)
        self.assertIn("已复核", text)

    def test_grade_sheet_has_criteria_columns(self):
        ws = self.wb["成绩表"]
        header_text = " ".join(
            str(c.value) for row in ws.iter_rows(min_row=1, max_row=6) for c in row if c.value
        )
        self.assertIn("观点理解", header_text)
        self.assertIn("个人思考", header_text)
        self.assertIn("文字表达", header_text)
        self.assertIn("满分5", header_text)

    def test_grade_sheet_has_average_row(self):
        ws = self.wb["成绩表"]
        text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertIn("平均", text)

    def test_rubric_sheet_documents_the_rules(self):
        ws = self.wb["评分规则"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertIn("观点理解", text)
        self.assertIn("待改进", text)
        self.assertIn("课堂心得", text)
        self.assertIn("字数下限", text)

    def test_similarity_sheet_lists_the_copied_pair(self):
        ws = self.wb["雷同比对"]
        text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        self.assertIn("张三", text)
        self.assertIn("李四", text)
        self.assertIn("学生 A", text)

    def test_freeze_panes_and_filter_are_set(self):
        ws = self.wb["成绩表"]
        self.assertTrue(ws.freeze_panes)
        self.assertTrue(ws.auto_filter.ref)


class TestMockBanner(unittest.TestCase):
    def test_mock_export_is_marked(self):
        data = to_bytes(sample_results(), DEFAULT_RUBRIC, mock=True)
        wb = load(data)
        text = "\n".join(
            str(c.value) for row in wb["成绩表"].iter_rows() for c in row if c.value is not None
        )
        self.assertIn("演示分数", text)

    def test_live_export_is_not_marked(self):
        data = to_bytes(sample_results(), DEFAULT_RUBRIC, mock=False)
        wb = load(data)
        text = "\n".join(
            str(c.value) for row in wb["成绩表"].iter_rows() for c in row if c.value is not None
        )
        self.assertNotIn("演示分数", text)


class TestEmptyAndEdge(unittest.TestCase):
    def test_empty_results_still_produces_a_valid_file(self):
        data = to_bytes([], DEFAULT_RUBRIC)
        wb = load(data)
        self.assertEqual(wb.sheetnames, ["成绩表", "评分规则", "雷同比对"])

    def test_result_with_error_is_written_out(self):
        broken = GradeResult(name="张三", content="x", max_total=10, error="接口超时")
        wb = load(to_bytes([broken], DEFAULT_RUBRIC))
        text = "\n".join(
            str(c.value) for row in wb["成绩表"].iter_rows() for c in row if c.value is not None
        )
        self.assertIn("接口超时", text)

    def test_custom_rubric_columns_follow(self):
        custom = DEFAULT_RUBRIC.from_dict(DEFAULT_RUBRIC.to_dict())
        custom.criteria[0].max_score = 20
        custom.criteria.append(type(custom.criteria[0])("participation", "课堂参与", 5))
        results = grade_batch([{"name": "张三", "content": CONCRETE}], custom, mock_cfg())
        wb = load(to_bytes(results, custom))
        header = " ".join(
            str(c.value) for row in wb["成绩表"].iter_rows(min_row=1, max_row=6) for c in row if c.value
        )
        self.assertIn("课堂参与", header)
        self.assertIn("满分20", header)


class TestFilenames(unittest.TestCase):
    def test_filename_contains_class_and_time(self):
        when = datetime(2026, 5, 20, 14, 30)
        name = suggested_filename(DEFAULT_RUBRIC, "计科2201", when=when)
        self.assertTrue(name.endswith(".xlsx"))
        self.assertIn("计科2201", name)
        self.assertIn("20260520_1430", name)

    def test_filename_strips_illegal_characters(self):
        name = suggested_filename(DEFAULT_RUBRIC, '计科/2201:A*B?')
        for bad in '<>:"/\\|?*':
            self.assertNotIn(bad, name)

    def test_filename_without_class(self):
        name = suggested_filename(DEFAULT_RUBRIC, "")
        self.assertTrue(name.startswith("课堂心得成绩_"))


class TestSaveToDisk(unittest.TestCase):
    def test_save_writes_a_readable_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "成绩.xlsx"
            written = save(sample_results(), DEFAULT_RUBRIC, target, class_name="计科2201")
            self.assertTrue(written.exists())
            wb = load_workbook(written)
            self.assertEqual(wb.sheetnames, ["成绩表", "评分规则", "雷同比对"])

    def test_save_creates_parent_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "子目录" / "成绩.xlsx"
            save(sample_results(), DEFAULT_RUBRIC, target)
            self.assertTrue(target.exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
